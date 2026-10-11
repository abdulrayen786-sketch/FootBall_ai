# ============================================================
# AI-BASED FOOTBALL PLAYER TRACKING AND PERFORMANCE ANALYSIS
# Professional Streamlit Dashboard
# ============================================================

import os
import subprocess
from pathlib import Path

import cv2
import imageio_ffmpeg
import numpy as np
import pandas as pd
import streamlit as st
from ultralytics import YOLO

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Football Analytics Pro",
    page_icon="⚽",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# PROFESSIONAL CSS
# ============================================================

st.markdown(
    """
    <style>
    .stApp {
        background:
            radial-gradient(circle at 90% 0%, rgba(0, 180, 150, 0.14), transparent 30%),
            linear-gradient(135deg, #04111c 0%, #071c2b 55%, #04101a 100%);
    }

    .block-container {
        max-width: 1450px;
        padding-top: 1.5rem;
        padding-bottom: 2rem;
    }

    section[data-testid="stSidebar"] {
        background: linear-gradient(180deg, #03101b 0%, #061c2b 100%);
        border-right: 1px solid rgba(255,255,255,0.08);
    }

    div[data-testid="stMetric"] {
        background: linear-gradient(145deg, rgba(12,46,64,0.96), rgba(5,23,37,0.96));
        border: 1px solid rgba(0,220,170,0.22);
        border-radius: 16px;
        padding: 16px;
        box-shadow: 0 8px 25px rgba(0,0,0,0.25);
    }

    div[data-testid="stMetricLabel"] {
        color: #9db2c5 !important;
        font-weight: 600;
    }

    div[data-testid="stMetricValue"] {
        color: #ffffff !important;
        font-weight: 800;
    }

    .hero {
        padding: 28px;
        border-radius: 20px;
        background: linear-gradient(135deg, rgba(0,105,90,0.45), rgba(7,31,49,0.96));
        border: 1px solid rgba(0,220,170,0.25);
        margin-bottom: 24px;
        box-shadow: 0 15px 40px rgba(0,0,0,0.30);
    }

    .hero-status {
        display: inline-block;
        padding: 6px 13px;
        border-radius: 20px;
        background: rgba(0,220,170,0.12);
        color: #00e0aa;
        font-size: 0.82rem;
        font-weight: 700;
        margin-bottom: 10px;
    }

    .hero-title {
        font-size: 2.15rem;
        font-weight: 800;
        margin-bottom: 5px;
    }

    .hero-subtitle {
        color: #9eb5c7;
        font-size: 1rem;
    }

    .card {
        background: linear-gradient(145deg, rgba(10,39,56,0.95), rgba(5,20,32,0.95));
        border: 1px solid rgba(255,255,255,0.08);
        border-radius: 16px;
        padding: 18px;
        margin-bottom: 14px;
        box-shadow: 0 8px 30px rgba(0,0,0,0.22);
    }

    .section-title {
        font-size: 1.35rem;
        font-weight: 750;
        margin-top: 26px;
        margin-bottom: 14px;
    }

    .footer {
        text-align: center;
        color: #71879a;
        padding: 25px 0 10px;
        margin-top: 30px;
        border-top: 1px solid rgba(255,255,255,0.06);
    }

    video {
        border-radius: 14px;
        border: 1px solid rgba(255,255,255,0.10);
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# DIRECTORIES
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"

INPUT_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

MODEL_PATH = BASE_DIR / "yolo11n.pt"


# ============================================================
# LOAD MODEL
# ============================================================

@st.cache_resource
def load_model():
    return YOLO(str(MODEL_PATH))


if not MODEL_PATH.exists():
    st.error(
        "YOLO model not found. Put yolo11n.pt in the same folder as app.py."
    )
    st.stop()

try:
    model = load_model()
except Exception as e:
    st.error(f"Could not load YOLO model: {e}")
    st.stop()


# ============================================================
# HERO
# ============================================================

st.markdown(
    """
    <div class="hero">
        <div class="hero-status">● AI SYSTEM READY</div>
        <div class="hero-title">⚽ Football Analytics Pro</div>
        <div class="hero-subtitle">
            AI-Powered Player Tracking &amp; Performance Analysis using Computer Vision
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# SIDEBAR
# ============================================================

st.sidebar.markdown("## ⚽ Football Analytics")
st.sidebar.caption("AI Performance Platform")
st.sidebar.markdown("---")

st.sidebar.subheader("⚙️ Analysis Settings")

pixels_per_meter = st.sidebar.number_input(
    "Pixels per meter",
    min_value=1.0,
    max_value=1000.0,
    value=50.0,
    step=1.0,
    help="Calibration used to estimate distance and speed.",
)

st.sidebar.markdown("---")
st.sidebar.markdown(
    """
    **Modules**

    ⚽ Ball Detection  
    🏃 Player Tracking  
    📏 Distance Analysis  
    🚀 Speed Analysis  
    👥 Team Statistics  
    🤖 AI Summary  
    📄 PDF Report  
    📊 Excel Report
    """
)

st.sidebar.markdown("---")
st.sidebar.info(
    "Player IDs are computer-vision tracking IDs. "
    "Distance and speed are estimates and depend on calibration."
)
st.sidebar.caption("Football Analytics Pro • v1.0")


# ============================================================
# SESSION STATE
# ============================================================

defaults = {
    "analysis_complete": False,
    "video_path": None,
    "tracking_data": pd.DataFrame(),
    "ball_data": pd.DataFrame(),
    "stats": pd.DataFrame(),
    "fps": 30.0,
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ============================================================
# AI SUMMARY
# ============================================================

def generate_performance_summary(player):
    distance = float(player["distance_covered_m"])
    average_speed = float(player["average_speed_kmh"])
    max_speed = float(player["max_speed_kmh"])
    frames = int(player["frames_tracked"])

    if distance >= 1000:
        movement_level = "very high"
    elif distance >= 500:
        movement_level = "high"
    elif distance >= 200:
        movement_level = "moderate"
    else:
        movement_level = "limited"

    if max_speed >= 30:
        speed_level = "very high"
    elif max_speed >= 20:
        speed_level = "high"
    elif max_speed >= 10:
        speed_level = "moderate"
    else:
        speed_level = "low"

    if frames >= 1000:
        tracking_quality = "strong"
    elif frames >= 500:
        tracking_quality = "good"
    else:
        tracking_quality = "limited"

    return (
        f"**Tracking ID:** {int(player['player_id'])}\n\n"
        f"The player was tracked for **{frames:,} frames**. "
        f"The estimated distance covered was **{distance:.2f} m**, "
        f"with an estimated average speed of **{average_speed:.2f} km/h** "
        f"and an estimated maximum speed of **{max_speed:.2f} km/h**.\n\n"
        f"Detected movement level: **{movement_level}**. "
        f"Maximum-speed level: **{speed_level}**. "
        f"Tracking coverage: **{tracking_quality}**.\n\n"
        f"These observations are based on computer-vision tracking data "
        f"and should be interpreted as estimates."
    )


# ============================================================
# VIDEO ANALYSIS
# ============================================================

def analyze_video(input_path, pixels_per_meter):
    cap = cv2.VideoCapture(str(input_path))

    if not cap.isOpened():
        raise RuntimeError("Could not open the uploaded video.")

    fps = cap.get(cv2.CAP_PROP_FPS)
    if fps <= 0:
        fps = 30.0

    width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
    height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))

    temp_video = OUTPUT_DIR / "analyzed_football_temp.avi"
    final_video = OUTPUT_DIR / "analyzed_football.mp4"

    writer = cv2.VideoWriter(
        str(temp_video),
        cv2.VideoWriter_fourcc(*"XVID"),
        fps,
        (width, height),
    )

    player_records = []
    ball_records = []
    previous_positions = {}

    progress = st.progress(0)
    status = st.empty()

    frame_number = 0

    try:
        while True:
            ret, frame = cap.read()
            if not ret:
                break

            frame_number += 1

            results = model.track(
                frame,
                persist=True,
                tracker="bytetrack.yaml",
                classes=[0, 32],
                verbose=False,
            )

            annotated = frame.copy()

            if results:
                result = results[0]
                boxes = result.boxes

                if boxes is not None:
                    for box in boxes:
                        if box.id is None:
                            continue

                        track_id = int(box.id.item())
                        cls = int(box.cls.item())
                        confidence = float(box.conf.item())

                        x1, y1, x2, y2 = map(
                            int, box.xyxy[0].tolist()
                        )

                        cx = (x1 + x2) // 2
                        cy = (y1 + y2) // 2

                        if cls == 0:
                            player_records.append(
                                {
                                    "frame": frame_number,
                                    "player_id": track_id,
                                    "x": cx,
                                    "y": cy,
                                    "confidence": confidence,
                                }
                            )

                            previous_positions[track_id] = (cx, cy)

                            cv2.rectangle(
                                annotated,
                                (x1, y1),
                                (x2, y2),
                                (0, 255, 0),
                                2,
                            )

                            cv2.putText(
                                annotated,
                                f"Player {track_id}",
                                (x1, max(20, y1 - 10)),
                                cv2.FONT_HERSHEY_SIMPLEX,
                                0.6,
                                (0, 255, 0),
                                2,
                            )

                        elif cls == 32:
                            ball_records.append(
                                {
                                    "frame": frame_number,
                                    "ball_id": track_id,
                                    "x": cx,
                                    "y": cy,
                                    "confidence": confidence,
                                }
                            )

                            cv2.circle(
                                annotated,
                                (cx, cy),
                                6,
                                (0, 0, 255),
                                -1,
                            )

                            cv2.putText(
                                annotated,
                                "BALL",
                                (cx + 8, cy),
                                cv2.FONT_HERSHEY_SIMPLEX,
                                0.6,
                                (0, 0, 255),
                                2,
                            )

            writer.write(annotated)

            if total_frames > 0:
                ratio = min(frame_number / total_frames, 1.0)
                progress.progress(ratio)
                status.write(
                    f"Processing frame {frame_number:,} / {total_frames:,}"
                )

    finally:
        cap.release()
        writer.release()
        progress.empty()
        status.empty()

    player_df = pd.DataFrame(player_records)
    ball_df = pd.DataFrame(ball_records)

    statistics = []

    if not player_df.empty:
        for player_id, group in player_df.groupby("player_id"):
            group = group.sort_values("frame")

            x = group["x"].to_numpy(dtype=float)
            y = group["y"].to_numpy(dtype=float)

            if len(group) > 1:
                dx = np.diff(x)
                dy = np.diff(y)
                movement_pixels = np.sqrt(dx ** 2 + dy ** 2)

                total_pixels = float(movement_pixels.sum())
                distance_m = total_pixels / pixels_per_meter

                duration_seconds = len(group) / fps
                average_speed_ms = (
                    distance_m / duration_seconds
                    if duration_seconds > 0
                    else 0.0
                )
                average_speed_kmh = average_speed_ms * 3.6

                instantaneous_speed_kmh = (
                    movement_pixels * fps / pixels_per_meter * 3.6
                )
                max_speed_kmh = (
                    float(instantaneous_speed_kmh.max())
                    if len(instantaneous_speed_kmh)
                    else 0.0
                )
            else:
                total_pixels = 0.0
                distance_m = 0.0
                average_speed_kmh = 0.0
                max_speed_kmh = 0.0

            statistics.append(
                {
                    "player_id": int(player_id),
                    "frames_tracked": int(len(group)),
                    "total_movement_pixels": total_pixels,
                    "distance_covered_m": distance_m,
                    "average_speed_kmh": average_speed_kmh,
                    "max_speed_kmh": max_speed_kmh,
                }
            )

    stats_df = pd.DataFrame(statistics)

    player_df.to_csv(
        OUTPUT_DIR / "player_tracking_data.csv",
        index=False,
    )
    ball_df.to_csv(
        OUTPUT_DIR / "ball_tracking_data.csv",
        index=False,
    )
    stats_df.to_csv(
        OUTPUT_DIR / "player_statistics.csv",
        index=False,
    )

    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()

    command = [
        ffmpeg,
        "-y",
        "-i",
        str(temp_video),
        "-c:v",
        "libx264",
        "-pix_fmt",
        "yuv420p",
        "-movflags",
        "+faststart",
        str(final_video),
    ]

    subprocess.run(
        command,
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )

    if temp_video.exists():
        temp_video.unlink()

    return final_video, player_df, ball_df, stats_df, fps


# ============================================================
# PDF REPORT
# ============================================================

def generate_pdf_report(stats, ball_data, output_path, pixels_per_meter):
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=35,
        leftMargin=35,
        topMargin=35,
        bottomMargin=35,
    )

    styles = getSampleStyleSheet()
    story = []

    story.append(
        Paragraph(
            "AI Football Performance Analysis Report",
            styles["Title"],
        )
    )
    story.append(Spacer(1, 10))
    story.append(
        Paragraph(
            "AI-Based Football Player Tracking and Performance Analysis "
            "Using Computer Vision",
            styles["Normal"],
        )
    )
    story.append(Spacer(1, 20))

    players = int(stats["player_id"].nunique()) if not stats.empty else 0
    total_distance = float(stats["distance_covered_m"].sum()) if not stats.empty else 0
    mean_speed = float(stats["average_speed_kmh"].mean()) if not stats.empty else 0
    max_speed = float(stats["max_speed_kmh"].max()) if not stats.empty else 0

    summary = [
        ["Metric", "Value"],
        ["Unique Player IDs", str(players)],
        ["Ball Detection Records", f"{len(ball_data):,}"],
        ["Sum of Tracked Distance", f"{total_distance:.2f} m"],
        ["Mean Player Average Speed", f"{mean_speed:.2f} km/h"],
        ["Highest Estimated Speed", f"{max_speed:.2f} km/h"],
        ["Calibration", f"{pixels_per_meter:.2f} pixels/m"],
    ]

    table = Table(summary, colWidths=[270, 170])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#087f68")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]
        )
    )

    story.append(table)
    story.append(Spacer(1, 25))
    story.append(Paragraph("Player Statistics", styles["Heading2"]))
    story.append(Spacer(1, 10))

    data = [["Player", "Frames", "Distance", "Avg Speed", "Max Speed"]]

    if not stats.empty:
        for _, row in stats.iterrows():
            data.append(
                [
                    str(int(row["player_id"])),
                    str(int(row["frames_tracked"])),
                    f'{row["distance_covered_m"]:.2f} m',
                    f'{row["average_speed_kmh"]:.2f}',
                    f'{row["max_speed_kmh"]:.2f}',
                ]
            )

    player_table = Table(
        data,
        repeatRows=1,
        colWidths=[70, 75, 105, 100, 100],
    )

    player_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#087f68")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )

    story.append(player_table)
    story.append(Spacer(1, 25))
    story.append(Paragraph("Methodology and Limitations", styles["Heading2"]))
    story.append(Spacer(1, 8))
    story.append(
        Paragraph(
            "YOLO detects players and the football. ByteTrack maintains "
            "tracking IDs between frames. Distance and speed are estimated "
            "using the supplied pixels-per-meter calibration. Camera "
            "perspective, missed detections, occlusion and tracking ID "
            "changes can affect the measurements.",
            styles["Normal"],
        )
    )

    doc.build(story)


# ============================================================
# EXCEL REPORT
# ============================================================

def generate_excel_report(
    stats,
    ball_data,
    tracking_data,
    output_path,
    pixels_per_meter,
):
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"

    summary["A1"] = "AI Football Performance Analysis Report"
    summary["A1"].font = Font(bold=True, size=16)

    summary["A3"] = "Metric"
    summary["B3"] = "Value"
    summary["A3"].font = Font(bold=True)
    summary["B3"].font = Font(bold=True)

    players = int(stats["player_id"].nunique()) if not stats.empty else 0
    total_distance = float(stats["distance_covered_m"].sum()) if not stats.empty else 0
    mean_speed = float(stats["average_speed_kmh"].mean()) if not stats.empty else 0
    max_speed = float(stats["max_speed_kmh"].max()) if not stats.empty else 0

    rows = [
        ("Unique Player IDs", players),
        ("Ball Detection Records", len(ball_data)),
        ("Sum of Tracked Distance (m)", total_distance),
        ("Mean Player Average Speed (km/h)", mean_speed),
        ("Highest Estimated Speed (km/h)", max_speed),
        ("Pixels per Meter", pixels_per_meter),
    ]

    for r, (name, value) in enumerate(rows, start=4):
        summary.cell(r, 1, name)
        summary.cell(r, 2, value)

    player_sheet = wb.create_sheet("Player Statistics")

    if not stats.empty:
        for c, name in enumerate(stats.columns, start=1):
            cell = player_sheet.cell(1, c, name)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for r, row in enumerate(stats.itertuples(index=False), start=2):
            for c, value in enumerate(row, start=1):
                player_sheet.cell(r, c, value)

    ball_sheet = wb.create_sheet("Ball Detection")

    if not ball_data.empty:
        for c, name in enumerate(ball_data.columns, start=1):
            cell = ball_sheet.cell(1, c, name)
            cell.font = Font(bold=True)

        for r, row in enumerate(ball_data.itertuples(index=False), start=2):
            for c, value in enumerate(row, start=1):
                ball_sheet.cell(r, c, value)

    tracking_sheet = wb.create_sheet("Player Tracking")

    if not tracking_data.empty:
        for c, name in enumerate(tracking_data.columns, start=1):
            cell = tracking_sheet.cell(1, c, name)
            cell.font = Font(bold=True)

        for r, row in enumerate(tracking_data.itertuples(index=False), start=2):
            for c, value in enumerate(row, start=1):
                tracking_sheet.cell(r, c, value)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        for column_cells in ws.columns:
            max_length = 0
            col_number = column_cells[0].column

            for cell in column_cells:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass

            ws.column_dimensions[get_column_letter(col_number)].width = min(
                max_length + 3, 35
            )

    wb.save(output_path)


# ============================================================
# UPLOAD SECTION
# ============================================================

st.markdown(
    '<div class="section-title">🎥 Upload Match Video</div>',
    unsafe_allow_html=True,
)

uploaded_file = st.file_uploader(
    "Choose a football video",
    type=["mp4", "avi", "mov", "mkv"],
)

if uploaded_file is not None:
    input_path = INPUT_DIR / uploaded_file.name

    with open(input_path, "wb") as f:
        f.write(uploaded_file.getbuffer())

    st.success(f"Video uploaded: {uploaded_file.name}")

    if st.button(
        "🚀 START FOOTBALL ANALYSIS",
        type="primary",
        use_container_width=True,
    ):
        try:
            with st.spinner("AI is analyzing the football video..."):
                (
                    final_video,
                    tracking_data,
                    ball_data,
                    stats,
                    fps,
                ) = analyze_video(
                    input_path,
                    pixels_per_meter,
                )

            st.session_state.analysis_complete = True
            st.session_state.video_path = str(final_video)
            st.session_state.tracking_data = tracking_data
            st.session_state.ball_data = ball_data
            st.session_state.stats = stats
            st.session_state.fps = fps

            st.success("✅ Football analysis completed successfully!")
            st.rerun()

        except Exception as e:
            st.error(f"Analysis failed: {e}")


# ============================================================
# RESULTS
# ============================================================

if st.session_state.analysis_complete:
    tracking_data = st.session_state.tracking_data
    ball_data = st.session_state.ball_data
    stats = st.session_state.stats
    video_path = st.session_state.video_path
    fps = st.session_state.fps

    # --------------------------------------------------------
    # OVERVIEW
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">📊 Match Overview</div>',
        unsafe_allow_html=True,
    )

    total_players = int(stats["player_id"].nunique()) if not stats.empty else 0
    total_records = len(tracking_data)
    ball_count = len(ball_data)
    total_distance = float(stats["distance_covered_m"].sum()) if not stats.empty else 0

    c1, c2, c3, c4 = st.columns(4)

    c1.metric("👥 Players Tracked", total_players)
    c2.metric("📍 Tracking Records", f"{total_records:,}")
    c3.metric("⚽ Ball Detections", f"{ball_count:,}")
    c4.metric("📏 Tracked Distance", f"{total_distance:.1f} m")

    # --------------------------------------------------------
    # VIDEO
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">🎥 Analyzed Match Video</div>',
        unsafe_allow_html=True,
    )

    if video_path and os.path.exists(video_path):
        st.video(video_path)
    else:
        st.warning("Analyzed video file was not found.")

    # --------------------------------------------------------
    # PLAYER ANALYSIS
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">🏃 Player Performance</div>',
        unsafe_allow_html=True,
    )

    if not stats.empty:
        player_ids = stats["player_id"].astype(int).tolist()

        selected_player = st.selectbox(
            "Select Player Tracking ID",
            player_ids,
        )

        row = stats[stats["player_id"] == selected_player].iloc[0]

        p1, p2, p3, p4 = st.columns(4)

        p1.metric("🆔 Tracking ID", int(row["player_id"]))
        p2.metric("🎞️ Frames", f'{int(row["frames_tracked"]):,}')
        p3.metric("📏 Distance", f'{row["distance_covered_m"]:.2f} m')
        p4.metric("🚀 Max Speed", f'{row["max_speed_kmh"]:.2f} km/h')

        player_data = tracking_data[
            tracking_data["player_id"] == selected_player
        ].copy()

        left, right = st.columns(2)

        with left:
            st.markdown("#### 📍 Player Movement")
            if not player_data.empty:
                movement = player_data[["x", "y"]].reset_index(drop=True)
                st.line_chart(movement)

        with right:
            st.markdown("#### 🚀 Speed Information")
            st.metric(
                "Average Speed",
                f'{row["average_speed_kmh"]:.2f} km/h',
            )
            st.metric(
                "Maximum Speed",
                f'{row["max_speed_kmh"]:.2f} km/h',
            )
            st.metric(
                "Tracking Coverage",
                f'{int(row["frames_tracked"]):,} frames',
            )

        st.markdown("#### 🤖 AI-Generated Performance Summary")
        st.info(generate_performance_summary(row))

        if len(player_data) > 1:
            x = player_data["x"].to_numpy(dtype=float)
            y = player_data["y"].to_numpy(dtype=float)

            movement_pixels = np.sqrt(
                np.diff(x) ** 2 + np.diff(y) ** 2
            )

            speed_kmh = (
                movement_pixels
                * fps
                / pixels_per_meter
                * 3.6
            )

            speed_df = pd.DataFrame(
                {
                    "Frame": player_data["frame"].iloc[1:].to_numpy(),
                    "Speed (km/h)": speed_kmh,
                }
            )

            st.markdown("#### 🚀 Speed Over Time")
            st.line_chart(speed_df.set_index("Frame"))

    else:
        st.info("No player statistics were generated.")

    # --------------------------------------------------------
    # TEAM STATISTICS
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">👥 Team Statistics</div>',
        unsafe_allow_html=True,
    )

    if not stats.empty:
        team_distance = float(stats["distance_covered_m"].sum())
        team_average_speed = float(stats["average_speed_kmh"].mean())
        highest_speed = float(stats["max_speed_kmh"].max())

        speed_row = stats.loc[stats["max_speed_kmh"].idxmax()]
        speed_player = int(speed_row["player_id"])

        t1, t2, t3, t4 = st.columns(4)

        t1.metric("👥 Players", int(stats["player_id"].nunique()))
        t2.metric("📏 Total Tracked Distance", f"{team_distance:.2f} m")
        t3.metric("🚶 Mean Average Speed", f"{team_average_speed:.2f} km/h")
        t4.metric("🚀 Highest Estimated Speed", f"{highest_speed:.2f} km/h")

        st.caption(
            f"Tracking ID {speed_player} recorded the highest estimated speed. "
            "The team distance is the sum of tracked-player distances."
        )

        a, b = st.columns(2)

        with a:
            st.markdown("#### 📏 Distance by Player")
            chart = stats[
                ["player_id", "distance_covered_m"]
            ].set_index("player_id")
            st.bar_chart(chart)

        with b:
            st.markdown("#### 🚀 Maximum Speed by Player")
            chart = stats[
                ["player_id", "max_speed_kmh"]
            ].set_index("player_id")
            st.bar_chart(chart)

    # --------------------------------------------------------
    # BALL DETECTION
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">⚽ Ball Detection</div>',
        unsafe_allow_html=True,
    )

    if not ball_data.empty:
        st.dataframe(
            ball_data,
            use_container_width=True,
            height=280,
        )
    else:
        st.info("No ball detections were recorded.")

    # --------------------------------------------------------
    # FULL PLAYER TABLE
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">📈 Complete Player Statistics</div>',
        unsafe_allow_html=True,
    )

    if not stats.empty:
        st.dataframe(
            stats,
            use_container_width=True,
            height=330,
        )

    # --------------------------------------------------------
    # REPORTS
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">📄 Reports & Export</div>',
        unsafe_allow_html=True,
    )

    pdf_path = OUTPUT_DIR / "football_performance_report.pdf"
    excel_path = OUTPUT_DIR / "football_performance_report.xlsx"

    r1, r2 = st.columns(2)

    with r1:
        st.markdown(
            """
            <div class="card">
                <h3>📄 PDF Match Report</h3>
                Player statistics, match summary and methodology.
            </div>
            """,
            unsafe_allow_html=True,
        )

        if st.button(
            "📄 Generate PDF Report",
            use_container_width=True,
        ):
            try:
                generate_pdf_report(
                    stats,
                    ball_data,
                    pdf_path,
                    pixels_per_meter,
                )
                st.success("PDF report generated.")
            except Exception as e:
                st.error(f"PDF generation failed: {e}")

        if pdf_path.exists():
            with open(pdf_path, "rb") as f:
                st.download_button(
                    "⬇️ Download PDF",
                    f,
                    file_name="football_performance_report.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                )

    with r2:
        st.markdown(
            """
            <div class="card">
                <h3>📊 Excel Data Report</h3>
                Player statistics, ball detections and tracking data.
            </div>
            """,
            unsafe_allow_html=True,
        )

        if st.button(
            "📊 Generate Excel Report",
            use_container_width=True,
        ):
            try:
                generate_excel_report(
                    stats,
                    ball_data,
                    tracking_data,
                    excel_path,
                    pixels_per_meter,
                )
                st.success("Excel report generated.")
            except Exception as e:
                st.error(f"Excel generation failed: {e}")

        if excel_path.exists():
            with open(excel_path, "rb") as f:
                st.download_button(
                    "⬇️ Download Excel",
                    f,
                    file_name="football_performance_report.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )

    # --------------------------------------------------------
    # METHODOLOGY
    # --------------------------------------------------------

    st.markdown(
        '<div class="section-title">⚙️ Methodology & Accuracy</div>',
        unsafe_allow_html=True,
    )

    st.info(
        "YOLO detects players and the football. ByteTrack maintains "
        "tracking IDs between frames. Distance and speed are estimated "
        "using the configured pixels-per-meter calibration."
    )

    st.warning(
        "Measurements are estimates. Camera perspective, camera motion, "
        "missed detections, occlusion and tracking-ID changes can affect "
        "the results. Tracking IDs are not guaranteed to represent "
        "specific real-world players."
    )

    # --------------------------------------------------------
    # FOOTER
    # --------------------------------------------------------

    st.markdown(
        """
        <div class="footer">
            ⚽ <b>Football Analytics Pro</b><br>
            AI-Based Football Player Tracking and Performance Analysis<br>
            Python • YOLO • OpenCV • ByteTrack • Streamlit
        </div>
        """,
        unsafe_allow_html=True,
    )

else:
    # ========================================================
    # EMPTY STATE
    # ========================================================

    st.markdown(
        """
        <div class="card" style="text-align:center; padding:45px;">
            <div style="font-size:60px;">⚽</div>
            <h2>Ready to Analyze Your Match</h2>
            <p style="color:#9eb5c7;">
                Upload a football video above to start AI-powered
                player and ball tracking.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

